import math
from dotenv import load_dotenv
from langchain_core.tools import tool
from langchain_tavily import TavilySearch
from database import save_memory, search_memory
from rag import retrieve_from_rag

from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
import smtplib
from tavily import TavilyClient
import os
import requests
from pathlib import Path

from datetime import datetime

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER




load_dotenv()

BASE_DIR = Path(__file__).resolve().parent

load_dotenv(BASE_DIR / ".env")

FAISS_DIR = BASE_DIR / "faiss-db"
SQLITE_DB = BASE_DIR / "chatbot.db"


GENERATED_DIR = Path("generated")
GENERATED_DIR.mkdir(exist_ok=True)

CURRENT_THREAD_ID = "default"




def set_current_thread_id(thread_id: str):
    global CURRENT_THREAD_ID
    CURRENT_THREAD_ID = thread_id

tavily_api_key = os.getenv("TAVILY_API_KEY")

tavily_client = (
    TavilyClient(api_key=tavily_api_key)
    if tavily_api_key
    else None
)

@tool
def search_web(query: str) -> str:
    """
    Search the web using Tavily for current
    or up-to-date information.
    """
    if tavily_client is None:
        return "Tavily API key is not configured."

    try:
        response = tavily_client.search(query=query)
        return str(response)
    except Exception as e:
        return f"Search Error: {str(e)}"

@tool
def calculator(expression: str) -> str:
    """
    Useful for simple math calculations.
    Input should be a valid math expression.
    Example: 2 + 2, math.sqrt(16), 10 * 5
    """
    try:
        allowed = {
            "math": math,
            "abs": abs,
            "round": round,
            "min": min,
            "max": max,
            "sum": sum
        }
        result = eval(expression, {"__builtins__": {}}, allowed)
        return str(result)
    except Exception as e:
        return f"Calculation error: {str(e)}"

@tool
def search_uploaded_documents(query: str) -> str:
    """
    Search uploaded documents for relevant information.
    Use this when the user asks about uploaded PDFs, DOCX, TXT, notes, files, or documents.
    """
    return retrieve_from_rag(
        query=query,
        thread_id=CURRENT_THREAD_ID
    )

@tool
def remember_this(memory: str) -> str:
    """
    Save an important user preference or fact into long-term memory.
    Use this when the user asks you to remember something.
    """
    return save_memory(
        thread_id=CURRENT_THREAD_ID,
        memory=memory
    )

@tool
def recall_memory(query: str) -> str:
    """
    Recall saved long-term memories about the user or this conversation.
    """
    return search_memory(
        thread_id=CURRENT_THREAD_ID,
        query=query
    )

@tool
def get_stock_price(symbol: str) -> dict:
    """
    Fetch the latest stock quote for a stock symbol.
    """
    api_key = os.getenv("TWELVE_DATA_API_KEY")

    if not api_key:
        return {"error": "TWELVE_DATA_API_KEY is not configured."}

    try:
        response = requests.get(
            "https://api.twelvedata.com/quote",
            params={
                "symbol": symbol.upper().strip(),
                "apikey": api_key,
            },
            timeout=10,
        )
        response.raise_for_status()
        return response.json()
    except Exception as e:
        return {"error": str(e)}

@tool
def send_email(
    to: str,
    subject: str,
    body: str
) -> str:
    """
    Send an email using the configured Gmail account.
    """
    try:
        sender_email = os.getenv("GMAIL_ADDRESS")
        app_password = os.getenv("GMAIL_APP_PASSWORD")

        if not sender_email or not app_password:
            return "Email configuration is missing. Please check GMAIL_ADDRESS and GMAIL_APP_PASSWORD in your .env file."

        message = MIMEMultipart()
        message["From"] = sender_email
        message["To"] = to
        message["Subject"] = subject
        message.attach(MIMEText(body, "plain"))

        with smtplib.SMTP("smtp.gmail.com", 587, timeout=20) as server:
            server.starttls()
            server.login(sender_email, app_password)
            server.send_message(message)

        return f"Email successfully sent to {to}"

    except Exception as e:
        return f"Email Error: {str(e)}"

@tool
def get_weather(city: str) -> dict:
    """
    Get the current weather for a specific city.
    """
    api_key = os.getenv("WEATHER_API_KEY")

    if not api_key:
        return {"error": "WEATHER_API_KEY is not configured."}

    try:
        response = requests.get(
            "https://api.weatherapi.com/v1/current.json",
            params={
                "key": api_key,
                "q": city,
            },
            timeout=10,
        )
        response.raise_for_status()
        return response.json()
    except Exception as e:
        return {"error": str(e)}

    

def clean_document_content(text: str) -> str:
    """
    Remove common Markdown formatting from generated
    document content while keeping the actual text readable.
    """

    import re

    if not text:
        return ""

    text = str(text)

    # Remove heading markers:
    # ### Heading
    # ## Heading
    # # Heading
    text = re.sub(
        r"^\s*#{1,6}\s*",
        "",
        text,
        flags=re.MULTILINE,
    )

    # Remove bold:
    # **text**
    text = text.replace("**", "")

    # Remove italic Markdown:
    # *text*
    text = re.sub(
        r"(?<!\*)\*([^*\n]+)\*(?!\*)",
        r"\1",
        text,
    )

    # Remove inline code:
    # `code`
    text = re.sub(
        r"`([^`]+)`",
        r"\1",
        text,
    )

    # Convert Markdown links:
    # [Google](https://google.com)
    text = re.sub(
        r"\[([^\]]+)\]\([^)]+\)",
        r"\1",
        text,
    )

    # Remove horizontal rules
    text = re.sub(
        r"^\s*([-*_]){3,}\s*$",
        "",
        text,
        flags=re.MULTILINE,
    )

    # Convert Markdown bullets to clean bullets
    text = re.sub(
        r"^\s*[-*+]\s+",
        "• ",
        text,
        flags=re.MULTILINE,
    )

    # Clean excessive blank lines
    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    return text.strip()






    
@tool
def create_document(
    title: str,
    content: str,
    file_format: str = "docx",
) -> str:
    """
    Create a clean DOCX or PDF document.

    Markdown formatting is automatically removed
    so the generated document looks professional.
    """

    import json

    file_format = file_format.lower().strip()

    if file_format not in ["docx", "pdf"]:
        return json.dumps({
            "success": False,
            "error": "file_format must be 'docx' or 'pdf'."
        })

    # --------------------------------------------------------
    # CLEAN MARKDOWN
    # --------------------------------------------------------

    clean_title = clean_document_content(title)
    clean_content = clean_document_content(content)

    # --------------------------------------------------------
    # TIMESTAMP
    # --------------------------------------------------------

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    # --------------------------------------------------------
    # SAFE FILE NAME
    # --------------------------------------------------------

    safe_title = "".join(
        c if c.isalnum() or c in (" ", "_", "-") else "_"
        for c in clean_title
    ).strip()

    if not safe_title:
        safe_title = "generated_document"

    filename = (
        f"{safe_title}_{timestamp}.{file_format}"
    )

    output_path = GENERATED_DIR / filename

    try:

        # ====================================================
        # DOCX
        # ====================================================

        if file_format == "docx":

            document = Document()

            # Title
            document.add_heading(
                clean_title,
                level=1,
            )

            # Content
            paragraphs = clean_content.split("\n")

            for paragraph in paragraphs:

                paragraph = paragraph.strip()

                if not paragraph:
                    continue

                # ------------------------------------------------
                # Handle bullet points
                # ------------------------------------------------

                if paragraph.startswith("• "):

                    document.add_paragraph(
                        paragraph[2:].strip(),
                        style="List Bullet",
                    )

                else:

                    document.add_paragraph(
                        paragraph
                    )

            document.save(
                output_path
            )

        # ====================================================
        # PDF
        # ====================================================

        elif file_format == "pdf":

            styles = getSampleStyleSheet()

            title_style = styles["Title"]
            title_style.alignment = TA_CENTER

            body_style = styles["BodyText"]
            body_style.leading = 16

            pdf = SimpleDocTemplate(
                str(output_path),
                pagesize=A4,
                rightMargin=50,
                leftMargin=50,
                topMargin=50,
                bottomMargin=50,
            )

            story = []

            # ------------------------------------------------
            # Title
            # ------------------------------------------------

            story.append(
                Paragraph(
                    clean_title,
                    title_style,
                )
            )

            story.append(
                Spacer(1, 20)
            )

            # ------------------------------------------------
            # Content
            # ------------------------------------------------

            paragraphs = clean_content.split("\n")

            for paragraph in paragraphs:

                paragraph = paragraph.strip()

                if not paragraph:
                    continue

                # Escape HTML characters
                paragraph = (
                    paragraph
                    .replace("&", "&amp;")
                    .replace("<", "&lt;")
                    .replace(">", "&gt;")
                )

                # ------------------------------------------------
                # Bullet
                # ------------------------------------------------

                if paragraph.startswith("• "):

                    paragraph = (
                        "• "
                        + paragraph[2:].strip()
                    )

                story.append(
                    Paragraph(
                        paragraph,
                        body_style,
                    )
                )

                story.append(
                    Spacer(1, 10)
                )

            pdf.build(story)

        # ====================================================
        # VERIFY FILE
        # ====================================================

        if not output_path.exists():

            return json.dumps({
                "success": False,
                "error": "Document was not created."
            })

        # ====================================================
        # RETURN DOCUMENT INFORMATION
        # ====================================================

        return json.dumps({
            "success": True,
            "type": file_format,
            "filename": filename,
            "download_url": (
                f"/generated/{filename}"
            )
        })

    except Exception as e:

        return json.dumps({
            "success": False,
            "error": str(e)
        })


tools = [
    calculator,
    search_uploaded_documents,
    remember_this,
    recall_memory,
    search_web,
    get_stock_price,
    send_email,
    get_weather,
    create_document
]